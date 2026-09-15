"""课表：作息设置、课程 CRUD、三种导入解析与一键生成任务。"""
import csv
import io
import json
import re
from datetime import date, timedelta

from fastapi import APIRouter, Depends, HTTPException, Query, Request
from sqlalchemy import select
from sqlalchemy.orm import Session

from ..database import get_db
from ..llm import chat_completion, extract_error_message
from ..models import Course, Plan, Task, TimetableSettings, User
from ..schemas import (
    CourseBulkRequest,
    CourseCreate,
    CourseImportPreview,
    CourseOut,
    CourseUpdate,
    FileImportPreview,
    GenerateTasksRequest,
    GenerateTasksResult,
    RowMappingRequest,
    TextImportRequest,
    TimetableSettingsOut,
    TimetableSettingsUpdate,
)
from ..security import get_current_user
from .plans import _resolve_llm

router = APIRouter(prefix="/api/timetable", tags=["timetable"])

MAX_COURSES = 300
MAX_IMPORT_ROWS = 500
MAX_WEEKS = 30
MAX_FILE_BYTES = 5 * 1024 * 1024
MAX_WARNINGS = 20
TIMETABLE_PLAN_TITLE = "课表"
DEFAULT_MINUTES = 90

# 默认作息：12 节，每节 45 分钟（可在课表页修改）
DEFAULT_PERIODS: list[dict] = [
    {"index": 1, "start": "08:00", "end": "08:45"},
    {"index": 2, "start": "08:55", "end": "09:40"},
    {"index": 3, "start": "10:00", "end": "10:45"},
    {"index": 4, "start": "10:55", "end": "11:40"},
    {"index": 5, "start": "14:00", "end": "14:45"},
    {"index": 6, "start": "14:55", "end": "15:40"},
    {"index": 7, "start": "16:00", "end": "16:45"},
    {"index": 8, "start": "16:55", "end": "17:40"},
    {"index": 9, "start": "19:00", "end": "19:45"},
    {"index": 10, "start": "19:55", "end": "20:40"},
    {"index": 11, "start": "20:50", "end": "21:35"},
    {"index": 12, "start": "21:45", "end": "22:30"},
]

_WEEKDAY_CN = {"一": 1, "二": 2, "三": 3, "四": 4, "五": 5, "六": 6, "日": 7, "天": 7}
_WEEKDAY_EN = (
    ("mon", 1),
    ("tue", 2),
    ("wed", 3),
    ("thu", 4),
    ("fri", 5),
    ("sat", 6),
    ("sun", 7),
)
_TIME_RE = re.compile(r"(\d{1,2})\s*[:：]\s*(\d{2})")
_RANGE_RE = re.compile(r"(\d+)\s*(?:-|~|—|－|至|到)\s*(\d+)")
_TIME_ONLY_RE = re.compile(r"^\d{1,2}:\d{2}$")

FIELD_KEYWORDS: dict[str, tuple[str, ...]] = {
    "name": ("课程名", "课程名称", "教学班", "课程", "科目", "名称"),
    "teacher": ("任课教师", "授课教师", "教师", "老师"),
    "location": ("上课地点", "上课教室", "地点", "教室", "场地"),
    "weekday": ("星期几", "星期", "周几"),
    "periods": ("节次", "节数", "节"),
    "weeks": ("起止周", "上课周次", "周次", "周数"),
    "time": ("上课时间", "时间", "时段"),
}


# ---------------- 基础工具 ----------------

def _normalize(value) -> str:
    """全角转半角并去空白，便于统一解析。"""
    text = "" if value is None else str(value)
    text = text.translate(
        str.maketrans("０１２３４５６７８９－～，、：　（）", "0123456789-~,,: ()")
    )
    return text.strip()


def _to_minutes(value: str) -> int:
    match = _TIME_RE.search(_normalize(value))
    if not match:
        return -1
    return int(match.group(1)) * 60 + int(match.group(2))


# ---------------- 解析：星期 / 节次 / 周次 ----------------

def parse_weekday(value) -> int | None:
    """「周一 / 星期一 / 一 / 1 / Mon」→ 1-7（周日为 7）。"""
    text = _normalize(value).lower()
    if not text:
        return None
    for key, day in _WEEKDAY_EN:
        if key in text:
            return day
    core = re.sub(r"(星期|周|礼拜|週|week|day)", "", text)
    for char in core:
        if char in _WEEKDAY_CN:
            return _WEEKDAY_CN[char]
    digits = re.findall(r"\d+", core)
    if digits:
        number = int(digits[0])
        if number == 0:
            return 7
        if 1 <= number <= 7:
            return number
    return None


def looks_like_time(value) -> bool:
    return bool(_TIME_RE.search(_normalize(value)))


def parse_periods(value) -> tuple[int, int] | None:
    """「1-2 / 1,2 / 第1-2节 / 3」→ (起, 止)。"""
    text = _normalize(value)
    if not text:
        return None
    numbers = [int(n) for n in re.findall(r"\d+", re.sub(r"[^\d\-~,;/]", "", text))]
    if not numbers:
        return None
    if len(numbers) == 1:
        return (numbers[0], numbers[0])
    return (min(numbers), max(numbers))


def parse_weeks(value) -> list[int] | None:
    """「1-16周 / 单周 / 双周 / 1,3,5 / 1-8,10-16」→ 周次数组。"""
    if isinstance(value, list):
        numbers = set()
        for item in value:
            try:
                numbers.add(int(item))
            except (TypeError, ValueError):
                continue
        weeks = sorted(n for n in numbers if 1 <= n <= MAX_WEEKS)
        return weeks or None
    text = _normalize(value)
    if not text:
        return None
    odd = "单" in text
    even = "双" in text
    numbers: set[int] = set()
    for start, end in _RANGE_RE.findall(text):
        low, high = sorted((int(start), int(end)))
        numbers.update(range(low, high + 1))
    rest = _RANGE_RE.sub(" ", text)
    for number in re.findall(r"\d+", rest):
        numbers.add(int(number))
    if not numbers and (odd or even):
        numbers = set(range(1, MAX_WEEKS + 1))
    if odd and not even:
        numbers = {n for n in numbers if n % 2 == 1}
    elif even and not odd:
        numbers = {n for n in numbers if n % 2 == 0}
    weeks = sorted(n for n in numbers if 1 <= n <= MAX_WEEKS)
    return weeks or None


def _slot_containing(periods: list[dict], minutes: int) -> int | None:
    for slot in periods:
        if _to_minutes(slot["start"]) <= minutes <= _to_minutes(slot["end"]):
            return int(slot["index"])
    fallback = [
        int(slot["index"])
        for slot in periods
        if _to_minutes(slot["start"]) <= minutes
    ]
    return max(fallback) if fallback else None


def parse_time_to_periods(value, periods: list[dict]) -> tuple[int, int] | None:
    """把「08:00-09:40」按作息表反查成节次区间。"""
    text = _normalize(value)
    matches = _TIME_RE.findall(text)
    if not matches:
        return None
    start_minutes = int(matches[0][0]) * 60 + int(matches[0][1])
    end_minutes = (
        int(matches[-1][0]) * 60 + int(matches[-1][1])
        if len(matches) > 1
        else start_minutes
    )
    start_slot = _slot_containing(periods, start_minutes)
    end_slot = _slot_containing(periods, end_minutes)
    if start_slot is None:
        return None
    return (start_slot, end_slot or start_slot)


def resolve_periods(period_text, time_text, periods: list[dict]) -> tuple[int, int] | None:
    """优先按节次文本解析；否则按时间文本反查节次。"""
    if period_text not in (None, "") and not looks_like_time(period_text):
        parsed = parse_periods(period_text)
        if parsed is not None:
            return parsed
    for candidate in (period_text, time_text):
        if candidate in (None, ""):
            continue
        parsed = parse_time_to_periods(candidate, periods)
        if parsed is not None:
            return parsed
    return None


# ---------------- 设置 ----------------

def _get_or_create_settings(db: Session, user_id: int) -> TimetableSettings:
    row = db.scalar(
        select(TimetableSettings).where(TimetableSettings.user_id == user_id)
    )
    if row is None:
        row = TimetableSettings(user_id=user_id)
        db.add(row)
        db.commit()
        db.refresh(row)
    return row


def _periods_of(row: TimetableSettings) -> list[dict]:
    raw = row.periods
    if isinstance(raw, str):
        try:
            raw = json.loads(raw)
        except (TypeError, ValueError):
            raw = None
    periods: list[dict] = []
    if isinstance(raw, list):
        for item in raw:
            if not isinstance(item, dict):
                continue
            try:
                index = int(item.get("index"))
            except (TypeError, ValueError):
                continue
            start = _normalize(item.get("start"))
            end = _normalize(item.get("end"))
            if not _TIME_ONLY_RE.match(start) or not _TIME_ONLY_RE.match(end):
                continue
            if _to_minutes(start) >= _to_minutes(end):
                continue
            periods.append({"index": index, "start": start, "end": end})
    if not periods:
        return [dict(slot) for slot in DEFAULT_PERIODS]
    return sorted(periods, key=lambda slot: slot["index"])


def _current_week(term_start: date | None, total_weeks: int) -> int | None:
    if term_start is None:
        return None
    offset = (date.today() - term_start).days
    if offset < 0:
        return None
    week = offset // 7 + 1
    return week if 1 <= week <= total_weeks else None


def _merged_settings(row: TimetableSettings) -> dict:
    total_weeks = row.total_weeks if 1 <= (row.total_weeks or 0) <= MAX_WEEKS else 16
    return {
        "term_start": row.term_start,
        "total_weeks": total_weeks,
        "periods": _periods_of(row),
        "current_week": _current_week(row.term_start, total_weeks),
    }


@router.get("/settings", response_model=TimetableSettingsOut)
def get_settings(
    user: User = Depends(get_current_user), db: Session = Depends(get_db)
):
    return _merged_settings(_get_or_create_settings(db, user.id))


@router.put("/settings", response_model=TimetableSettingsOut)
def update_settings(
    payload: TimetableSettingsUpdate,
    user: User = Depends(get_current_user),
    db: Session = Depends(get_db),
):
    row = _get_or_create_settings(db, user.id)
    changes = payload.model_dump(exclude_unset=True)
    if "term_start" in changes:
        row.term_start = changes["term_start"]
    if "total_weeks" in changes and changes["total_weeks"] is not None:
        row.total_weeks = changes["total_weeks"]
    if "periods" in changes and changes["periods"] is not None:
        slots = sorted(
            (slot if isinstance(slot, dict) else slot.model_dump() for slot in changes["periods"]),
            key=lambda slot: slot["index"],
        )
        indexes = [int(slot["index"]) for slot in slots]
        if indexes != list(range(1, len(slots) + 1)):
            raise HTTPException(status_code=400, detail="节次需从 1 开始且连续")
        for slot in slots:
            if _to_minutes(slot["start"]) >= _to_minutes(slot["end"]):
                raise HTTPException(
                    status_code=400,
                    detail=f"第 {slot['index']} 节的结束时间需晚于开始时间",
                )
        row.periods = [
            {"index": int(slot["index"]), "start": _normalize(slot["start"]), "end": _normalize(slot["end"])}
            for slot in slots
        ]
    db.commit()
    db.refresh(row)
    return _merged_settings(row)


# ---------------- 课程 CRUD ----------------

def _normalize_course(data: dict) -> dict:
    start = int(data.get("start_period") or 1)
    end = int(data.get("end_period") or start)
    if end < start:
        raise HTTPException(status_code=400, detail="结束节次不能早于开始节次")
    data["start_period"] = max(1, min(30, start))
    data["end_period"] = max(1, min(30, end))
    weeks = data.get("weeks")
    if weeks is not None:
        cleaned = sorted({int(w) for w in weeks if 1 <= int(w) <= MAX_WEEKS})
        data["weeks"] = cleaned or None
    return data


def _get_owned_course(course_id: int, user_id: int, db: Session) -> Course:
    course = db.scalar(
        select(Course).where(Course.id == course_id, Course.user_id == user_id)
    )
    if course is None:
        raise HTTPException(status_code=404, detail="课程不存在")
    return course


@router.get("/courses", response_model=list[CourseOut])
def list_courses(
    user: User = Depends(get_current_user), db: Session = Depends(get_db)
):
    return db.scalars(
        select(Course)
        .where(Course.user_id == user.id)
        .order_by(Course.weekday, Course.start_period, Course.id)
    ).all()


@router.post("/courses", response_model=CourseOut, status_code=201)
def create_course(
    payload: CourseCreate,
    user: User = Depends(get_current_user),
    db: Session = Depends(get_db),
):
    total = len(db.scalars(select(Course.id).where(Course.user_id == user.id)).all())
    if total >= MAX_COURSES:
        raise HTTPException(status_code=400, detail=f"课程数量已达上限（{MAX_COURSES} 条）")
    data = _normalize_course(payload.model_dump())
    course = Course(**data, user_id=user.id)
    db.add(course)
    db.commit()
    db.refresh(course)
    return course


@router.patch("/courses/{course_id}", response_model=CourseOut)
def update_course(
    course_id: int,
    payload: CourseUpdate,
    user: User = Depends(get_current_user),
    db: Session = Depends(get_db),
):
    course = _get_owned_course(course_id, user.id, db)
    changes = payload.model_dump(exclude_unset=True)
    merged = {
        "start_period": changes.get("start_period", course.start_period),
        "end_period": changes.get("end_period", course.end_period),
        "weeks": changes.get("weeks", course.weeks),
    }
    _normalize_course(merged)
    changes["start_period"] = merged["start_period"]
    changes["end_period"] = merged["end_period"]
    if "weeks" in changes or "start_period" in changes or "end_period" in changes:
        changes["weeks"] = merged["weeks"]
    for key, value in changes.items():
        setattr(course, key, value)
    db.commit()
    db.refresh(course)
    return course


@router.delete("/courses/{course_id}", status_code=204)
def delete_course(
    course_id: int,
    user: User = Depends(get_current_user),
    db: Session = Depends(get_db),
):
    course = _get_owned_course(course_id, user.id, db)
    db.delete(course)
    db.commit()


@router.post("/courses/bulk", response_model=list[CourseOut], status_code=201)
def bulk_save_courses(
    payload: CourseBulkRequest,
    user: User = Depends(get_current_user),
    db: Session = Depends(get_db),
):
    if payload.replace:
        for course in db.scalars(select(Course).where(Course.user_id == user.id)).all():
            db.delete(course)
        db.flush()
    existing = len(db.scalars(select(Course.id).where(Course.user_id == user.id)).all())
    if existing + len(payload.courses) > MAX_COURSES:
        raise HTTPException(
            status_code=400,
            detail=f"课程数量将超过上限（{MAX_COURSES} 条），请减少导入或先清空",
        )
    created: list[Course] = []
    for item in payload.courses:
        data = _normalize_course(item.model_dump())
        course = Course(**data, user_id=user.id)
        db.add(course)
        created.append(course)
    db.commit()
    for course in created:
        db.refresh(course)
    return created


# ---------------- 导入解析 ----------------

def _cell(row: list, column) -> str:
    try:
        index = int(column)
    except (TypeError, ValueError):
        return ""
    if index < 0 or index >= len(row):
        return ""
    return _normalize(row[index])


def _coerce_course(item: dict, periods: list[dict]) -> tuple[CourseCreate | None, list[str]]:
    """把任意来源的一条记录规整成课程；无法识别时返回原因。"""
    if not isinstance(item, dict):
        return None, ["跳过一条无法识别的记录"]
    name = _normalize(item.get("name") or item.get("title") or item.get("course"))[:100]
    if not name:
        return None, ["跳过一条缺少课程名的记录"]

    weekday_source = item.get("weekday")
    if weekday_source is None:
        weekday_source = item.get("week")
    weekday = parse_weekday(weekday_source)
    if weekday is None:
        return None, [f"「{name}」缺少可识别的星期，已跳过"]

    period_text = item.get("periods")
    if period_text in (None, ""):
        period_text = item.get("period")
    if period_text in (None, "") and item.get("start_period") is not None:
        period_text = f"{item.get('start_period')}-{item.get('end_period') or item.get('start_period')}"
    time_text = item.get("time") or item.get("start_time")
    parsed = resolve_periods(period_text, time_text, periods)
    if parsed is None:
        return None, [f"「{name}」缺少可识别的节次或时间，已跳过"]
    start_period, end_period = parsed
    start_period = max(1, min(30, start_period))
    end_period = max(start_period, min(30, end_period))

    course = CourseCreate(
        name=name,
        teacher=_normalize(item.get("teacher"))[:50],
        location=_normalize(item.get("location"))[:100],
        weekday=weekday,
        start_period=start_period,
        end_period=end_period,
        weeks=parse_weeks(item.get("weeks") or item.get("week_text") or item.get("week_range")),
        color=_normalize(item.get("color"))[:20],
        note=_normalize(item.get("note"))[:200],
    )
    return course, []


def _conflict_warnings(courses: list[CourseCreate]) -> list[str]:
    warnings: list[str] = []
    for index, first in enumerate(courses):
        for second in courses[index + 1 :]:
            if first.weekday != second.weekday:
                continue
            if first.end_period < second.start_period or second.end_period < first.start_period:
                continue
            weeks_a = set(first.weeks or [])
            weeks_b = set(second.weeks or [])
            if weeks_a and weeks_b and not (weeks_a & weeks_b):
                continue
            warnings.append(
                f"周{first.weekday} 第 {first.start_period}-{first.end_period} 节："
                f"「{first.name}」与「{second.name}」时间冲突"
            )
            if len(warnings) >= MAX_WARNINGS:
                return warnings
    return warnings


def _suggest_mapping(headers: list[str]) -> dict[str, int]:
    mapping: dict[str, int] = {}
    for field, keywords in FIELD_KEYWORDS.items():
        for index, header in enumerate(headers):
            text = _normalize(header)
            if text and any(keyword in text for keyword in keywords):
                mapping[field] = index
                break
    if "weekday" not in mapping:
        for index, header in enumerate(headers):
            text = _normalize(header)
            if "周" in text and not any(k in text for k in ("周次", "周数", "起止")):
                mapping["weekday"] = index
                break
    return mapping


def _preview_from_courses(courses: list[CourseCreate]) -> CourseImportPreview:
    warnings: list[str] = []
    if len(courses) > MAX_COURSES:
        warnings.append(f"解析出 {len(courses)} 条，只保留前 {MAX_COURSES} 条")
        courses = courses[:MAX_COURSES]
    warnings.extend(_conflict_warnings(courses))
    return CourseImportPreview(courses=courses, warnings=warnings)


def _read_csv_rows(body: bytes) -> list[list[str]]:
    text = None
    for encoding in ("utf-8-sig", "gb18030", "utf-16", "latin-1"):
        try:
            text = body.decode(encoding)
            break
        except (UnicodeDecodeError, LookupError):
            continue
    if text is None:
        raise HTTPException(status_code=400, detail="无法识别 CSV 文件编码，请另存为 UTF-8")
    return [[_normalize(cell) for cell in row[:20]] for row in csv.reader(io.StringIO(text))]


def _read_xlsx_rows(body: bytes) -> list[list[str]]:
    try:
        from openpyxl import load_workbook
    except ImportError:  # pragma: no cover - 依赖缺失时给出明确提示
        raise HTTPException(
            status_code=400, detail="服务器缺少 openpyxl 依赖，无法解析 Excel，请改用 CSV"
        )
    try:
        workbook = load_workbook(io.BytesIO(body), read_only=True, data_only=True)
    except Exception:  # noqa: BLE001 - 文件损坏等一律按无法解析处理
        raise HTTPException(status_code=400, detail="Excel 文件无法解析，请确认是 .xlsx 格式")
    try:
        list_rows: list[list[str]] = []
        for row in workbook.active.iter_rows(values_only=True):
            list_rows.append([_normalize(cell) for cell in row[:20]])
            if len(list_rows) > MAX_IMPORT_ROWS + 1:
                break
    finally:
        workbook.close()
    return list_rows


@router.post("/import/parse-file", response_model=FileImportPreview)
async def parse_file(
    request: Request,
    filename: str = Query("", max_length=200),
    user: User = Depends(get_current_user),
    db: Session = Depends(get_db),
):
    """上传 xlsx/csv，返回表头、预览行与建议列映射（不落库）。"""
    body = await request.body()
    if not body:
        raise HTTPException(status_code=400, detail="文件内容为空")
    if len(body) > MAX_FILE_BYTES:
        raise HTTPException(status_code=400, detail="文件过大（上限 5MB）")
    name = filename.strip().lower()
    if name.endswith(".xls"):
        raise HTTPException(
            status_code=400, detail="暂不支持旧版 .xls，请另存为 .xlsx 或 .csv 后重试"
        )
    if name.endswith(".csv"):
        rows = _read_csv_rows(body)
    elif name.endswith(".xlsx"):
        rows = _read_xlsx_rows(body)
    else:
        raise HTTPException(status_code=400, detail="只支持 .xlsx 与 .csv 文件")

    rows = [row for row in rows if any(cell for cell in row)]
    if not rows:
        raise HTTPException(status_code=400, detail="文件里没有可解析的内容")
    headers = rows[0]
    data_rows = rows[1:]
    total = len(data_rows)
    truncated = total > MAX_IMPORT_ROWS
    return FileImportPreview(
        headers=headers,
        rows=data_rows[:MAX_IMPORT_ROWS],
        suggested_mapping=_suggest_mapping(headers),
        total_rows=total,
        truncated=truncated,
    )


@router.post("/import/parse-rows", response_model=CourseImportPreview)
def parse_rows(
    payload: RowMappingRequest,
    user: User = Depends(get_current_user),
    db: Session = Depends(get_db),
):
    """按列映射把表格行解析成课程（不落库）。"""
    periods = _merged_settings(_get_or_create_settings(db, user.id))["periods"]
    courses: list[CourseCreate] = []
    warnings: list[str] = []
    for index, row in enumerate(payload.rows, start=1):
        item = {field: _cell(row, column) for field, column in payload.mapping.items()}
        course, notes = _coerce_course(item, periods)
        if course is None:
            if any(cell for cell in row):
                warnings.extend(f"第 {index} 行：{note}" for note in notes)
            continue
        courses.append(course)
    preview = _preview_from_courses(courses)
    preview.warnings = (warnings + preview.warnings)[: MAX_WARNINGS * 2]
    return preview


@router.post("/import/parse-text", response_model=CourseImportPreview)
def parse_text(
    payload: TextImportRequest,
    user: User = Depends(get_current_user),
    db: Session = Depends(get_db),
):
    """粘贴课表文本，交给 AI 解析成课程（不落库）。"""
    resolved = _resolve_llm(user.id, db)
    if resolved is None:
        raise HTTPException(
            status_code=400,
            detail="尚未配置 AI 服务：请先在「AI 设置」中配置 API，或由服务器设置 LLM_API_KEY 环境变量",
        )
    periods = _merged_settings(_get_or_create_settings(db, user.id))["periods"]
    prompt = (
        "请把下面的课表文本解析成结构化课程列表。\n"
        "规则：weekday 用 1-7 表示（1=周一，7=周日）；start_period/end_period 用阿拉伯数字表示第几节；"
        "weeks 是生效周次数组（如第 1-16 周为 [1,2,...,16]，单周只保留奇数周）。\n"
        "只返回 JSON，格式："
        '{"courses":[{"name":"课程名","teacher":"教师","location":"教室",'
        '"weekday":1,"start_period":1,"end_period":2,"weeks":[1,2,3]}]}\n'
        "课表文本：\n"
        f"{payload.text}"
    )
    try:
        content = chat_completion(
            resolved["base_url"],
            resolved["model"],
            resolved["api_key"],
            [{"role": "user", "content": prompt}],
        )
    except Exception as exc:  # noqa: BLE001 - 统一转成用户可读错误
        raise HTTPException(
            status_code=502, detail=f"AI 调用失败：{extract_error_message(exc)}"
        )

    data = _json_from_text(content)
    items = data.get("courses") if isinstance(data, dict) else data
    if not isinstance(items, list):
        raise HTTPException(status_code=502, detail="AI 返回内容无法解析，请重试或改用手动录入")

    courses: list[CourseCreate] = []
    warnings: list[str] = []
    for item in items:
        course, notes = _coerce_course(item, periods)
        if course is None:
            warnings.extend(notes)
            continue
        courses.append(course)
    if not courses:
        raise HTTPException(status_code=502, detail="AI 没有解析出可用课程，请检查文本或改用手动录入")
    preview = _preview_from_courses(courses)
    preview.warnings = (warnings + preview.warnings)[: MAX_WARNINGS * 2]
    return preview


def _json_from_text(content: str):
    text = (content or "").strip()
    if text.startswith("```"):
        text = re.sub(r"^```[a-zA-Z]*\n?", "", text)
        text = re.sub(r"\n?```$", "", text).strip()
    try:
        return json.loads(text)
    except (TypeError, ValueError):
        pass
    match = re.search(r"[\[{].*[\]}]", text, re.S)
    if not match:
        return None
    try:
        return json.loads(match.group(0))
    except (TypeError, ValueError):
        return None


# ---------------- 一键生成任务 ----------------

def _course_dates(course: Course, term_start: date, total_weeks: int) -> list[date]:
    weeks = course.weeks or list(range(1, total_weeks + 1))
    dates = set()
    for week in weeks:
        if not 1 <= week <= total_weeks:
            continue
        dates.add(term_start + timedelta(days=(week - 1) * 7 + (course.weekday - 1)))
    return sorted(dates)


def _slot_by_index(periods: list[dict], index: int) -> dict | None:
    for slot in periods:
        if int(slot["index"]) == int(index):
            return slot
    return None


def _minutes_for(course: Course, periods: list[dict]) -> int:
    start_slot = _slot_by_index(periods, course.start_period)
    end_slot = _slot_by_index(periods, course.end_period)
    if start_slot is None or end_slot is None:
        return DEFAULT_MINUTES
    delta = _to_minutes(end_slot["end"]) - _to_minutes(start_slot["start"])
    if delta <= 0:
        return DEFAULT_MINUTES
    return max(1, min(600, delta))


def _resolve_plan(db: Session, user: User, plan_id: int | None) -> Plan:
    if plan_id is not None:
        plan = db.scalar(
            select(Plan).where(Plan.id == plan_id, Plan.user_id == user.id)
        )
        if plan is None:
            raise HTTPException(status_code=404, detail="计划不存在")
        return plan
    plan = db.scalar(
        select(Plan).where(
            Plan.user_id == user.id,
            Plan.title == TIMETABLE_PLAN_TITLE,
            Plan.parent_id.is_(None),
        )
    )
    if plan is None:
        plan = Plan(
            user_id=user.id,
            title=TIMETABLE_PLAN_TITLE,
            description="由课表一键生成的上课任务",
        )
        db.add(plan)
        db.commit()
        db.refresh(plan)
    return plan


def _generate_for(
    db: Session, user: User, courses: list[Course], plan_id: int | None
) -> GenerateTasksResult:
    settings = _merged_settings(_get_or_create_settings(db, user.id))
    term_start = settings["term_start"]
    if term_start is None:
        raise HTTPException(
            status_code=400,
            detail="请先在课表页设置「学期第 1 周周一」，再生成任务",
        )
    plan = _resolve_plan(db, user, plan_id)
    total_weeks = settings["total_weeks"]
    periods = settings["periods"]

    created: list[Task] = []
    skipped = 0
    for course in courses:
        minutes = _minutes_for(course, periods)
        title = (
            f"{course.name}（{course.location}）" if course.location else course.name
        )[:200]
        for day in _course_dates(course, term_start, total_weeks):
            exists = db.scalar(
                select(Task.id).where(
                    Task.user_id == user.id,
                    Task.title == title,
                    Task.due_date == day,
                )
            )
            if exists is not None:
                skipped += 1
                continue
            task = Task(
                user_id=user.id,
                plan_id=plan.id,
                title=title,
                subject=course.name,
                estimated_minutes=minutes,
                due_date=day,
            )
            db.add(task)
            created.append(task)
    db.commit()
    for task in created:
        db.refresh(task)
    return GenerateTasksResult(
        created=len(created),
        skipped=skipped,
        plan_id=plan.id,
        tasks=created,
        message=f"生成 {len(created)} 条上课任务，跳过已存在 {skipped} 条",
    )


@router.post("/courses/{course_id}/generate-tasks", response_model=GenerateTasksResult, status_code=201)
def generate_course_tasks(
    course_id: int,
    payload: GenerateTasksRequest,
    user: User = Depends(get_current_user),
    db: Session = Depends(get_db),
):
    course = _get_owned_course(course_id, user.id, db)
    return _generate_for(db, user, [course], payload.plan_id)


@router.post("/generate-tasks", response_model=GenerateTasksResult, status_code=201)
def generate_all_tasks(
    payload: GenerateTasksRequest,
    user: User = Depends(get_current_user),
    db: Session = Depends(get_db),
):
    courses = list(
        db.scalars(select(Course).where(Course.user_id == user.id)).all()
    )
    if not courses:
        raise HTTPException(status_code=400, detail="还没有课程，请先导入或手动添加课表")
    return _generate_for(db, user, courses, payload.plan_id)
